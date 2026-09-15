# -*- coding: utf-8 -*-
"""采集流程的端到端测试（需要 Playwright 浏览器内核）。"""
import json

import pytest
from openpyxl import load_workbook

from web_agent import WebAgent
from web_agent.extract import collect

pytestmark = pytest.mark.usefixtures("need_browser")


def test_collect_notice_list_fields(base_url):
    res = collect(base_url, "/notices", fields=["标题", "发布时间", "链接"])
    assert res["kind"] == "table"
    assert res["missing_fields"] == []
    assert res["count"] == 6
    assert all(r["标题"] for r in res["records"])
    assert res["records"][0]["链接"].startswith(base_url)


def test_collect_products_cards(base_url):
    res = collect(base_url, "/products", fields=["名称", "型号", "价格", "供应商"])
    assert res["kind"] == "list"
    assert res["missing_fields"] == []
    assert res["count"] == 8
    first = res["records"][0]
    assert first["名称"] == "高性能计算服务器"
    assert first["型号"] == "HS-7020"
    assert first["价格"] == "128000 元"


def test_collect_detail_page_as_single_record(base_url):
    res = collect(base_url, "/notices/N-01", kind="item",
                  fields=["标题", "发布单位", "发布时间", "正文", "链接"])
    assert res["count"] == 1
    record = res["records"][0]
    assert record["发布单位"] == "网络空间安全学院"
    assert record["发布时间"] == "2026-09-01"
    assert record["链接"].endswith("/files/N-01.pdf")


def test_all_columns_when_fields_omitted(base_url):
    res = collect(base_url, "/notices")
    assert list(res["records"][0].keys()) == ["编号", "标题", "发布单位", "发布时间", "分类", "详情"]


def test_multi_page_collect_with_dedup(base_url):
    res = WebAgent(base_url).collect_pages("/notices", "all", ["编号", "标题", "链接"], "编号")
    assert res["pages"] == [1, 2, 3]
    assert res["count"] == 12
    assert res["duplicates_removed"] == 1


def test_natural_language_export_csv(base_url, tmp_path):
    out = tmp_path / "notices.csv"
    result = WebAgent(base_url).run("把前两页公告的标题、发布时间和链接导出成 CSV", out=str(out))
    assert result["ok"] is True
    assert result["plan"]["source"] == "rule"
    assert result["plan"]["format"] == "csv"
    assert result["count"] == 11
    assert out.exists() and out.read_text(encoding="utf-8-sig").startswith("标题")


def test_natural_language_export_excel_all_pages(base_url, tmp_path):
    out = tmp_path / "notices.xlsx"
    result = WebAgent(base_url).run("采集全部公告的标题、发布单位和发布时间，导出成 Excel", out=str(out))
    assert result["plan"]["pages"] == "all"
    assert result["count"] == 13
    sheet = load_workbook(out).active
    assert sheet.max_row == 14
    assert [c.value for c in sheet[1]] == ["标题", "发布单位", "发布时间"]


def test_natural_language_dedup_removes_duplicate(base_url, tmp_path):
    out = tmp_path / "dedup.json"
    result = WebAgent(base_url).run("采集前两页公告的编号和标题并去重", out=str(out))
    assert result["count"] == 10
    assert result["duplicates_removed"] == 1
    assert len(json.loads(out.read_text(encoding="utf-8"))) == 10
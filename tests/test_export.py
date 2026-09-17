# -*- coding: utf-8 -*-
"""导出功能的单元测试：JSON / CSV / Excel，以及写文件前的覆盖确认。"""
import io
import json

import pytest
from openpyxl import load_workbook

from web_agent.export import columns_of, export, to_csv_text, to_json_text

ROWS = [{"标题": "甲", "发布时间": "2026-09-01"},
        {"标题": "乙", "发布时间": "2026-09-02"}]


def test_columns_of_union_keeps_order():
    assert columns_of([{"a": 1}, {"b": 2, "a": 3}]) == ["a", "b"]


def test_to_json_text_is_utf8_readable():
    data = json.loads(to_json_text(ROWS))
    assert data[0]["标题"] == "甲"
    assert len(data) == 2


def test_to_csv_text_has_header_and_rows():
    lines = to_csv_text(ROWS).strip().splitlines()
    assert lines[0] == "标题,发布时间"
    assert lines[1].startswith("甲,")


def test_export_json_file(tmp_path):
    path = tmp_path / "out.json"
    info = export(ROWS, "json", str(path))
    assert info["count"] == 2
    assert info["columns"] == ["标题", "发布时间"]
    assert json.loads(path.read_text(encoding="utf-8"))[1]["标题"] == "乙"


def test_export_csv_file_is_excel_friendly(tmp_path):
    path = tmp_path / "out.csv"
    export(ROWS, "csv", str(path))
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")  # utf-8-sig


def test_export_xlsx_roundtrip(tmp_path):
    path = tmp_path / "out.xlsx"
    export(ROWS, "xlsx", str(path))
    sheet = load_workbook(path).active
    assert [c.value for c in sheet[1]] == ["标题", "发布时间"]
    assert sheet["A2"].value == "甲"
    assert sheet.max_row == 3


def test_export_rejects_unknown_format():
    with pytest.raises(ValueError):
        export(ROWS, "yaml", None)


def test_xlsx_requires_output_path():
    with pytest.raises(ValueError):
        export(ROWS, "xlsx", None)


def test_export_without_path_writes_nothing():
    info = export(ROWS, "json", None)
    assert info["written"] is False
    assert info["count"] == 2


# ---------------- 写文件前的覆盖确认（有风险操作的用户确认） ----------------

def _stub_confirm(answer: bool):
    """记录被询问过的路径，并固定返回 answer。"""
    calls = []

    def confirm(path):
        calls.append(path)
        return answer

    return calls, confirm


def test_existing_file_is_kept_when_user_declines(tmp_path):
    path = tmp_path / "out.json"
    path.write_text("旧内容", encoding="utf-8")
    calls, confirm = _stub_confirm(False)
    info = export(ROWS, "json", str(path), confirm=confirm)
    assert info["written"] is False
    assert "取消写入" in info["message"]
    assert calls == [str(path)]                          # 确实问过用户
    assert path.read_text(encoding="utf-8") == "旧内容"   # 原文件未被改动


def test_confirmed_overwrite_replaces_file(tmp_path):
    path = tmp_path / "out.json"
    path.write_text("旧内容", encoding="utf-8")
    calls, confirm = _stub_confirm(True)
    info = export(ROWS, "json", str(path), confirm=confirm)
    assert info["written"] is True
    assert calls == [str(path)]
    assert json.loads(path.read_text(encoding="utf-8"))[1]["标题"] == "乙"


def test_new_file_is_written_without_asking(tmp_path):
    calls, confirm = _stub_confirm(True)
    info = export(ROWS, "json", str(tmp_path / "brand_new.json"), confirm=confirm)
    assert info["written"] is True
    assert calls == []                                   # 不打扰用户


def test_overwrite_flag_skips_confirmation(tmp_path):
    path = tmp_path / "out.csv"
    path.write_text("旧内容", encoding="utf-8")
    calls, confirm = _stub_confirm(True)
    info = export(ROWS, "csv", str(path), overwrite=True, confirm=confirm)
    assert info["written"] is True
    assert calls == []
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_non_interactive_environment_declines_overwrite(tmp_path, monkeypatch):
    path = tmp_path / "out.json"
    path.write_text("旧内容", encoding="utf-8")
    monkeypatch.setattr("sys.stdin", io.StringIO(""))     # 模拟管道/自动化：无人应答
    info = export(ROWS, "json", str(path))
    assert info["written"] is False
    assert path.read_text(encoding="utf-8") == "旧内容"


def test_answer_with_byte_order_mark_still_confirms(tmp_path, monkeypatch):
    """从管道或带 BOM 的文件读入的应答不应影响确认判断。"""
    path = tmp_path / "out.json"
    path.write_text("旧内容", encoding="utf-8")
    monkeypatch.setattr("sys.stdin", io.StringIO("\ufeffy\n"))
    info = export(ROWS, "json", str(path))
    assert info["written"] is True


def test_xlsx_overwrite_also_asks(tmp_path):
    path = tmp_path / "out.xlsx"
    export(ROWS, "xlsx", str(path))
    before = path.read_bytes()
    calls, confirm = _stub_confirm(False)
    info = export(ROWS, "xlsx", str(path), confirm=confirm)
    assert info["written"] is False
    assert calls == [str(path)]
    assert path.read_bytes() == before
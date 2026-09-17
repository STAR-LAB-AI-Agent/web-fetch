# -*- coding: utf-8 -*-
"""采集结果导出：JSON / CSV / Excel。

三种格式的定位：
- JSON：保留完整的结构化数据，便于程序继续处理；
- CSV ：便于用表格软件查看，使用 utf-8-sig 编码以便 Excel 直接打开不乱码；
- XLSX：使用 openpyxl 写入，包含表头样式与列宽设置。

写文件属于有风险的操作：目标路径上已有同名文件时，默认先询问用户是否覆盖，
得到确认后才写入；用户拒绝或处于非交互环境时跳过写入，并在返回值中说明。
"""
from __future__ import annotations

import csv
import io
import json
import os
import sys
from typing import Mapping, Sequence

FORMATS = ("json", "csv", "xlsx")


def columns_of(rows: Sequence[Mapping[str, object]]) -> list:
    """按出现顺序汇总所有列名，容忍不同行字段不一致的情况。"""
    cols = []
    for row in rows:
        for key in row.keys():
            if key not in cols:
                cols.append(key)
    return cols


def to_json_text(rows: Sequence[Mapping[str, object]]) -> str:
    return json.dumps(list(rows), ensure_ascii=False, indent=2)


def to_csv_text(rows: Sequence[Mapping[str, object]]) -> str:
    cols = columns_of(rows)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({c: row.get(c, "") for c in cols})
    return buf.getvalue()


def write_xlsx(rows: Sequence[Mapping[str, object]], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    cols = columns_of(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "采集结果"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF1F6")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([str(row.get(c, "")) for c in cols])
    for i, col in enumerate(cols, start=1):
        width = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows]) if rows else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width * 1.4))
    wb.save(path)


def ask_overwrite(path: str) -> bool:
    """在终端询问用户是否覆盖已有文件。

    提示写到标准错误，避免污染命令行工具输出到标准输出的 JSON 结果；
    标准输入已经结束（管道、重定向等非交互场景）时视为拒绝覆盖。
    """
    sys.stderr.write("该文件已存在，是否覆盖？[y/N] %s " % os.path.abspath(path))
    sys.stderr.flush()
    line = sys.stdin.readline()
    sys.stderr.write("\n")
    # 从管道或带 BOM 的文件里读到的内容可能夹带字节序标记，先清掉再判断
    return line.replace("\ufeff", "").strip().lower() in ("y", "yes")


def export(rows: Sequence[Mapping[str, object]], fmt: str = "json", path: str | None = None, *,
           overwrite: bool = False, confirm=None) -> dict:
    """导出采集结果，返回 {format, path, count, columns, written[, message]}。

    - 未指定输出路径时只返回数据概况，不写任何文件；
    - 目标文件已存在且 overwrite 为假时，先用 confirm 询问用户，确认后才覆盖；
      用户拒绝（或非交互环境下无人应答）时跳过写入，written 为 False。
    """
    fmt = (fmt or "json").lower()
    if fmt not in FORMATS:
        raise ValueError("不支持的导出格式：%s（可选 %s）" % (fmt, " / ".join(FORMATS)))
    if fmt == "xlsx" and not path:
        raise ValueError("导出 Excel 必须指定输出文件路径")

    info = {
        "format": fmt,
        "path": os.path.abspath(path) if path else None,
        "count": len(rows),
        "columns": columns_of(rows),
        "written": bool(path),
    }

    if path and os.path.exists(path) and not overwrite:
        if not (confirm or ask_overwrite)(path):
            info["written"] = False
            info["message"] = "目标文件已存在，已按用户选择取消写入"
            return info

    if path:
        if fmt == "json":
            with open(path, "w", encoding="utf-8") as fp:
                fp.write(to_json_text(rows))
        elif fmt == "csv":
            with open(path, "w", encoding="utf-8-sig", newline="") as fp:
                fp.write(to_csv_text(rows))
        else:
            write_xlsx(rows, path)

    return info